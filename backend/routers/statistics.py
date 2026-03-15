"""
统计信息路由
"""

from fastapi import APIRouter, Depends, HTTPException, status, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, and_
from backend.database import get_db
from backend.models import Order, ServiceRecord, Supplier, User
from backend.schemas import (
    StatisticsResponse, StatisticsItem,
    UserStatisticsResponse, UserStatisticsItem, GroupStatisticsItem
)
from backend.auth import get_current_user, can_view_internal_price
from backend.utils import parse_order_content, calculate_order_totals
from backend.config import (
    USER_TYPE_ADMIN, USER_TYPE_NORMAL, USER_TYPE_SUPPLIER, USER_TYPE_STUDENT,
    ORDER_STATUS_CONFIRMED, SERVICE_STATUS_CONFIRMED, TAX_RATE
)
from backend.logger import logger
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
import io

router = APIRouter(prefix="/api/statistics", tags=["statistics"])


@router.get("/", response_model=StatisticsResponse)
async def get_statistics(
    start_date: str = Query(None, description="开始日期，格式：YYYY-MM-DD"),
    end_date: str = Query(None, description="结束日期，格式：YYYY-MM-DD"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    获取统计信息（课题组用户和管理员）
    按照供应商分组统计订单和服务记录

    Args:
        start_date: 开始日期（可选），格式：YYYY-MM-DD
        end_date: 结束日期（可选），格式：YYYY-MM-DD
        current_user: 当前登录用户
        db: 数据库会话

    Returns:
        StatisticsResponse: 统计信息响应

    Raises:
        HTTPException: 如果用户类型不允许

    使用样例:
        GET /api/statistics/
        GET /api/statistics/?start_date=2025-01-01&end_date=2025-12-31
    """
    # 只有课题组用户和管理员可以查看统计信息（普通用户不能查看）
    if current_user.user_type == USER_TYPE_SUPPLIER or current_user.user_type == USER_TYPE_STUDENT:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="供应商用户和普通用户不能查看统计信息"
        )
    
    can_view_internal = can_view_internal_price(current_user)

    # 构建订单查询（只统计确认状态的订单）
    order_query = db.query(Order).filter(Order.status == ORDER_STATUS_CONFIRMED)

    # 添加时间范围过滤
    if start_date:
        try:
            start_dt = datetime.strptime(start_date, "%Y-%m-%d")
            order_query = order_query.filter(Order.created_at >= start_dt)
        except ValueError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="开始日期格式错误，应为YYYY-MM-DD"
            )
    if end_date:
        try:
            # 结束日期包含当天，所以需要加23:59:59
            end_dt = datetime.strptime(end_date, "%Y-%m-%d").replace(hour=23, minute=59, second=59)
            order_query = order_query.filter(Order.created_at <= end_dt)
        except ValueError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="结束日期格式错误，应为YYYY-MM-DD"
            )
    if current_user.user_type == USER_TYPE_NORMAL:
        # 课题组用户统计自己的订单以及其管理学生的订单
        from sqlalchemy import or_
        managed_students = db.query(User.id).filter(User.manager_id == current_user.id).all()
        managed_student_ids = [s[0] for s in managed_students]
        if managed_student_ids:
            order_query = order_query.filter(
                or_(
                    Order.user_id == current_user.id,
                    Order.user_id.in_(managed_student_ids)
                )
            )
        else:
            order_query = order_query.filter(Order.user_id == current_user.id)

    # 构建服务记录查询（只统计确认状态的服务）
    service_query = db.query(ServiceRecord).filter(ServiceRecord.status == SERVICE_STATUS_CONFIRMED)

    # 添加时间范围过滤（服务记录）
    if start_date:
        start_dt = datetime.strptime(start_date, "%Y-%m-%d")
        service_query = service_query.filter(ServiceRecord.created_at >= start_dt)
    if end_date:
        end_dt = datetime.strptime(end_date, "%Y-%m-%d").replace(hour=23, minute=59, second=59)
        service_query = service_query.filter(ServiceRecord.created_at <= end_dt)
    if current_user.user_type == USER_TYPE_NORMAL:
        # 课题组用户统计自己的服务记录以及其管理学生的服务记录
        from sqlalchemy import or_
        managed_students = db.query(User.id).filter(User.manager_id == current_user.id).all()
        managed_student_ids = [s[0] for s in managed_students]
        if managed_student_ids:
            service_query = service_query.filter(
                or_(
                    ServiceRecord.user_id == current_user.id,
                    ServiceRecord.user_id.in_(managed_student_ids)
                )
            )
        else:
            service_query = service_query.filter(ServiceRecord.user_id == current_user.id)
    
    # 获取所有相关的供应商ID
    supplier_ids = set()
    orders = order_query.all()
    services = service_query.all()
    
    for order in orders:
        supplier_ids.add(order.supplier_id)
    for service in services:
        supplier_ids.add(service.supplier_id)
    
    # 按供应商统计
    statistics_items = []
    total_order_count = 0
    total_product_count = 0
    total_internal_price = 0.0
    total_tax_included_price = 0.0
    total_service_amount = 0.0
    
    for supplier_id in supplier_ids:
        supplier = db.query(Supplier).filter(Supplier.id == supplier_id).first()
        if not supplier:
            continue
        
        # 统计该供应商的订单
        supplier_orders = [o for o in orders if o.supplier_id == supplier_id]
        order_count = len(supplier_orders)
        product_count = 0
        supplier_internal_price = 0.0
        supplier_tax_included_price = 0.0
        
        for order in supplier_orders:
            items = parse_order_content(order.content)
            product_count += sum(item.get("quantity", 1) for item in items)
            totals = calculate_order_totals(items, include_internal=can_view_internal)
            if can_view_internal:
                supplier_internal_price += totals["total_internal_price"]
            supplier_tax_included_price += totals["total_tax_included_price"]
        
        # 统计该供应商的服务记录
        supplier_services = [s for s in services if s.supplier_id == supplier_id]
        supplier_service_amount = sum(s.amount for s in supplier_services)
        
        # 计算税额和结余
        # 总税额 = (订单总含税价格 - 订单总团购价格) * 13%
        supplier_tax = (supplier_tax_included_price - supplier_internal_price) * TAX_RATE
        # 总结余 = (订单总含税价格 - 订单总团购价格) - 总税额 - 总服务价格
        supplier_balance = (supplier_tax_included_price - supplier_internal_price) - supplier_tax - supplier_service_amount
        
        statistics_items.append(StatisticsItem(
            supplier_id=supplier.id,
            supplier_name=supplier.name,
            order_count=order_count,
            product_count=product_count,
            total_internal_price=supplier_internal_price,
            total_tax_included_price=supplier_tax_included_price,
            total_service_amount=supplier_service_amount,
            total_tax=supplier_tax,
            total_balance=supplier_balance
        ))
        
        # 累加总计
        total_order_count += order_count
        total_product_count += product_count
        total_internal_price += supplier_internal_price
        total_tax_included_price += supplier_tax_included_price
        total_service_amount += supplier_service_amount
    
    # 计算总计
    # 总税额 = (订单总含税价格 - 订单总团购价格) * 13%
    total_tax = (total_tax_included_price - total_internal_price) * TAX_RATE
    # 总结余 = (订单总含税价格 - 订单总团购价格) - 总税额 - 总服务价格
    total_balance = (total_tax_included_price - total_internal_price) - total_tax - total_service_amount
    
    total_item = StatisticsItem(
        supplier_id=0,
        supplier_name="总计",
        order_count=total_order_count,
        product_count=total_product_count,
        total_internal_price=total_internal_price,
        total_tax_included_price=total_tax_included_price,
        total_service_amount=total_service_amount,
        total_tax=total_tax,
        total_balance=total_balance
    )
    
    return StatisticsResponse(
        items=statistics_items,
        total=total_item
    )


@router.get("/by-user", response_model=UserStatisticsResponse)
async def get_statistics_by_user(
    start_date: str = Query(None, description="开始日期，格式：YYYY-MM-DD"),
    end_date: str = Query(None, description="结束日期，格式：YYYY-MM-DD"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    获取按用户分组的统计信息（课题组用户和管理员）

    返回两个表格的数据：
    1. 课题组分组统计（groups + group_total）
    2. 用户详细统计（details + detail_total）

    Args:
        start_date: 开始日期（可选），格式：YYYY-MM-DD
        end_date: 结束日期（可选），格式：YYYY-MM-DD
        current_user: 当前登录用户
        db: 数据库会话

    Returns:
        UserStatisticsResponse: 按用户统计的响应

    Raises:
        HTTPException: 如果用户类型不允许

    使用样例:
        GET /api/statistics/by-user
        GET /api/statistics/by-user?start_date=2025-01-01&end_date=2025-12-31
    """
    # 只有课题组用户和管理员可以查看统计信息（普通用户不能查看）
    if current_user.user_type == USER_TYPE_SUPPLIER or current_user.user_type == USER_TYPE_STUDENT:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="供应商用户和普通用户不能查看统计信息"
        )

    can_view_internal = can_view_internal_price(current_user)

    # 构建订单查询（只统计确认状态的订单）
    order_query = db.query(Order).filter(Order.status == ORDER_STATUS_CONFIRMED)

    # 添加时间范围过滤
    if start_date:
        try:
            start_dt = datetime.strptime(start_date, "%Y-%m-%d")
            order_query = order_query.filter(Order.created_at >= start_dt)
        except ValueError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="开始日期格式错误，应为YYYY-MM-DD"
            )
    if end_date:
        try:
            end_dt = datetime.strptime(end_date, "%Y-%m-%d").replace(hour=23, minute=59, second=59)
            order_query = order_query.filter(Order.created_at <= end_dt)
        except ValueError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="结束日期格式错误，应为YYYY-MM-DD"
            )

    if current_user.user_type == USER_TYPE_NORMAL:
        # 课题组用户统计自己的订单以及其管理学生的订单
        from sqlalchemy import or_
        managed_students = db.query(User.id).filter(User.manager_id == current_user.id).all()
        managed_student_ids = [s[0] for s in managed_students]
        if managed_student_ids:
            order_query = order_query.filter(
                or_(
                    Order.user_id == current_user.id,
                    Order.user_id.in_(managed_student_ids)
                )
            )
        else:
            order_query = order_query.filter(Order.user_id == current_user.id)

    # 构建服务记录查询（只统计确认状态的服务）
    service_query = db.query(ServiceRecord).filter(ServiceRecord.status == SERVICE_STATUS_CONFIRMED)

    # 添加时间范围过滤（服务记录）
    if start_date:
        start_dt = datetime.strptime(start_date, "%Y-%m-%d")
        service_query = service_query.filter(ServiceRecord.created_at >= start_dt)
    if end_date:
        end_dt = datetime.strptime(end_date, "%Y-%m-%d").replace(hour=23, minute=59, second=59)
        service_query = service_query.filter(ServiceRecord.created_at <= end_dt)

    if current_user.user_type == USER_TYPE_NORMAL:
        # 课题组用户统计自己的服务记录以及其管理学生的服务记录
        from sqlalchemy import or_
        managed_students = db.query(User.id).filter(User.manager_id == current_user.id).all()
        managed_student_ids = [s[0] for s in managed_students]
        if managed_student_ids:
            service_query = service_query.filter(
                or_(
                    ServiceRecord.user_id == current_user.id,
                    ServiceRecord.user_id.in_(managed_student_ids)
                )
            )
        else:
            service_query = service_query.filter(ServiceRecord.user_id == current_user.id)

    orders = order_query.all()
    services = service_query.all()

    # 获取所有涉及的用户
    user_ids = set()
    for order in orders:
        user_ids.add(order.user_id)
    for service in services:
        user_ids.add(service.user_id)

    users = db.query(User).filter(User.id.in_(user_ids)).all()
    user_dict = {u.id: u for u in users}

    # 构建用户到课题组的映射
    # 管理员订单 -> manager_id = 0
    # 课题组用户订单 -> manager_id = 自己的ID
    # 普通用户订单 -> manager_id = 其manager_id
    group_mapping = {}  # {user_id: group_manager_id}
    group_members = {}  # {group_manager_id: [user_ids]}

    for user_id in user_ids:
        user = user_dict.get(user_id)
        if not user:
            continue

        if user.user_type == USER_TYPE_ADMIN:
            group_mapping[user_id] = 0
        elif user.user_type == USER_TYPE_NORMAL:
            group_mapping[user_id] = user.id
        elif user.user_type == USER_TYPE_STUDENT:
            group_mapping[user_id] = user.manager_id if user.manager_id else 0

    # 构建组成员
    for user_id, group_id in group_mapping.items():
        if group_id not in group_members:
            group_members[group_id] = []
        group_members[group_id].append(user_id)

    # 如果是课题组用户，过滤掉管理员组
    if current_user.user_type == USER_TYPE_NORMAL and 0 in group_members:
        del group_members[0]

    # 按用户详细统计
    user_stats = {}
    for user_id in user_ids:
        user_orders = [o for o in orders if o.user_id == user_id]
        user_services = [s for s in services if s.user_id == user_id]

        user = user_dict.get(user_id)
        if not user:
            continue

        # 获取该用户的课题组信息
        group_id = group_mapping.get(user_id, 0)
        if group_id == 0:
            group_name = "管理员组"
        else:
            group_manager = user_dict.get(group_id)
            group_name = group_manager.username if group_manager else f"组{group_id}"

        # 计算该用户的统计数据
        order_count = len(user_orders)
        product_count = 0
        user_internal_price = 0.0
        user_tax_included_price = 0.0

        for order in user_orders:
            items = parse_order_content(order.content)
            product_count += sum(item.get("quantity", 1) for item in items)
            totals = calculate_order_totals(items, include_internal=can_view_internal)
            if can_view_internal:
                user_internal_price += totals["total_internal_price"]
            user_tax_included_price += totals["total_tax_included_price"]

        user_service_amount = sum(s.amount for s in user_services)
        user_tax = (user_tax_included_price - user_internal_price) * TAX_RATE
        user_balance = (user_tax_included_price - user_internal_price) - user_tax - user_service_amount

        user_stats[user_id] = UserStatisticsItem(
            user_id=user_id,
            username=user.username,
            group_id=group_id,
            group_name=group_name,
            order_count=order_count,
            product_count=product_count,
            total_internal_price=user_internal_price,
            total_tax_included_price=user_tax_included_price,
            total_service_amount=user_service_amount,
            total_tax=user_tax,
            total_balance=user_balance
        )

    # 按课题组分组统计
    group_stats = {}
    for group_id, member_ids in group_members.items():
        group_order_count = 0
        group_product_count = 0
        group_internal_price = 0.0
        group_tax_included_price = 0.0
        group_service_amount = 0.0

        for user_id in member_ids:
            stat = user_stats.get(user_id)
            if stat:
                group_order_count += stat.order_count
                group_product_count += stat.product_count
                group_internal_price += stat.total_internal_price
                group_tax_included_price += stat.total_tax_included_price
                group_service_amount += stat.total_service_amount

        group_tax = (group_tax_included_price - group_internal_price) * TAX_RATE
        group_balance = (group_tax_included_price - group_internal_price) - group_tax - group_service_amount

        # 确定组名
        if group_id == 0:
            group_name = "管理员组"
        else:
            group_manager = user_dict.get(group_id)
            group_name = group_manager.username if group_manager else f"组{group_id}"

        group_stats[group_id] = GroupStatisticsItem(
            manager_id=group_id,
            manager_name=group_name,
            user_count=len([uid for uid in member_ids if user_dict.get(uid) and user_dict.get(uid).user_type == USER_TYPE_STUDENT]),
            order_count=group_order_count,
            product_count=group_product_count,
            total_internal_price=group_internal_price,
            total_tax_included_price=group_tax_included_price,
            total_service_amount=group_service_amount,
            total_tax=group_tax,
            total_balance=group_balance
        )

    # 计算总计（确保两个总计使用相同的值）
    total_order_count = sum(s.order_count for s in user_stats.values())
    total_product_count = sum(s.product_count for s in user_stats.values())
    total_internal_price = sum(s.total_internal_price for s in user_stats.values())
    total_tax_included_price = sum(s.total_tax_included_price for s in user_stats.values())
    total_service_amount = sum(s.total_service_amount for s in user_stats.values())
    total_tax = (total_tax_included_price - total_internal_price) * TAX_RATE
    total_balance = (total_tax_included_price - total_internal_price) - total_tax - total_service_amount

    # 构建响应
    groups_list = list(group_stats.values())
    details_list = list(user_stats.values())

    return UserStatisticsResponse(
        groups=groups_list,
        group_total=GroupStatisticsItem(
            manager_id=0,
            manager_name="总计",
            user_count=sum(g.user_count for g in groups_list),
            order_count=total_order_count,
            product_count=total_product_count,
            total_internal_price=total_internal_price,
            total_tax_included_price=total_tax_included_price,
            total_service_amount=total_service_amount,
            total_tax=total_tax,
            total_balance=total_balance
        ),
        details=details_list,
        detail_total=UserStatisticsItem(
            user_id=0,
            username="总计",
            group_id=0,
            group_name="总计",
            order_count=total_order_count,
            product_count=total_product_count,
            total_internal_price=total_internal_price,
            total_tax_included_price=total_tax_included_price,
            total_service_amount=total_service_amount,
            total_tax=total_tax,
            total_balance=total_balance
        )
    )


@router.get("/export/supplier")
async def export_supplier_statistics(
    start_date: str = Query(None, description="开始日期，格式：YYYY-MM-DD"),
    end_date: str = Query(None, description="结束日期，格式：YYYY-MM-DD"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    导出按供应商统计的Excel文件

    Args:
        start_date: 开始日期（可选），格式：YYYY-MM-DD
        end_date: 结束日期（可选），格式：YYYY-MM-DD
        current_user: 当前登录用户
        db: 数据库会话

    Returns:
        StreamingResponse: Excel文件流

    Raises:
        HTTPException: 如果用户类型不允许

    使用样例:
        GET /api/statistics/export/supplier
        GET /api/statistics/export/supplier?start_date=2025-01-01&end_date=2025-12-31
    """
    # 调用统计接口获取数据
    stats_data = await get_statistics(start_date, end_date, current_user, db)

    # 创建Excel工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = "按供应商统计"

    # 设置标题样式
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")
    cell_alignment = Alignment(horizontal="right", vertical="center")
    currency_alignment = Alignment(horizontal="right", vertical="center")

    # 写入表头
    headers = ["供应商", "订单总数", "商品总数", "订单总团购价格", "订单总含税价格", "总服务价格", "总税额", "总结余"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # 写入数据
    for row, item in enumerate(stats_data.items, 2):
        ws.cell(row=row, column=1, value=item.supplier_name).alignment = cell_alignment
        ws.cell(row=row, column=2, value=item.order_count).alignment = cell_alignment
        ws.cell(row=row, column=3, value=item.product_count).alignment = cell_alignment
        ws.cell(row=row, column=4, value=item.total_internal_price).alignment = currency_alignment
        ws.cell(row=row, column=5, value=item.total_tax_included_price).alignment = currency_alignment
        ws.cell(row=row, column=6, value=item.total_service_amount).alignment = currency_alignment
        ws.cell(row=row, column=7, value=item.total_tax).alignment = currency_alignment
        ws.cell(row=row, column=8, value=item.total_balance).alignment = currency_alignment

    # 写入总计行
    total_row = len(stats_data.items) + 2
    total_cell = ws.cell(row=total_row, column=1, value=stats_data.total.supplier_name)
    total_cell.font = Font(bold=True)
    total_cell.alignment = cell_alignment

    for col in range(2, 9):
        cell = ws.cell(row=total_row, column=col)
        cell.font = Font(bold=True)
        cell.alignment = currency_alignment

    ws.cell(row=total_row, column=2, value=stats_data.total.order_count)
    ws.cell(row=total_row, column=3, value=stats_data.total.product_count)
    ws.cell(row=total_row, column=4, value=stats_data.total.total_internal_price)
    ws.cell(row=total_row, column=5, value=stats_data.total.total_tax_included_price)
    ws.cell(row=total_row, column=6, value=stats_data.total.total_service_amount)
    ws.cell(row=total_row, column=7, value=stats_data.total.total_tax)
    ws.cell(row=total_row, column=8, value=stats_data.total.total_balance)

    # 调整列宽
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 15

    # 保存到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    logger.info(
        f"按供应商统计导出成功: 用户={current_user.username}",
        extra={
            "user_id": current_user.id,
            "user_type": current_user.user_type,
            "start_date": start_date,
            "end_date": end_date
        }
    )

    from urllib.parse import quote
    filename = f"按供应商统计_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    encoded_filename = quote(filename, safe='')

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"
        }
    )


@router.get("/export/user")
async def export_user_statistics(
    start_date: str = Query(None, description="开始日期，格式：YYYY-MM-DD"),
    end_date: str = Query(None, description="结束日期，格式：YYYY-MM-DD"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    导出按用户统计的Excel文件

    Args:
        start_date: 开始日期（可选），格式：YYYY-MM-DD
        end_date: 结束日期（可选），格式：YYYY-MM-DD
        current_user: 当前登录用户
        db: 数据库会话

    Returns:
        StreamingResponse: Excel文件流

    Raises:
        HTTPException: 如果用户类型不允许

    使用样例:
        GET /api/statistics/export/user
        GET /api/statistics/export/user?start_date=2025-01-01&end_date=2025-12-31
    """
    # 调用统计接口获取数据
    stats_data = await get_statistics_by_user(start_date, end_date, current_user, db)

    # 创建Excel工作簿
    wb = Workbook()

    # 设置样式
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")
    cell_alignment = Alignment(horizontal="right", vertical="center")
    currency_alignment = Alignment(horizontal="right", vertical="center")
    group_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    # 第一个Sheet：课题组统计
    ws1 = wb.active
    ws1.title = "课题组统计"

    # 写入课题组统计表头
    headers1 = ["课题组", "用户数", "订单总数", "商品总数", "订单总团购价格", "订单总含税价格", "总服务价格", "总税额", "总结余"]
    for col, header in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # 写入课题组数据
    for row, item in enumerate(stats_data.groups, 2):
        ws1.cell(row=row, column=1, value=item.manager_name).alignment = cell_alignment
        ws1.cell(row=row, column=2, value=item.user_count).alignment = cell_alignment
        ws1.cell(row=row, column=3, value=item.order_count).alignment = cell_alignment
        ws1.cell(row=row, column=4, value=item.product_count).alignment = cell_alignment
        ws1.cell(row=row, column=5, value=item.total_internal_price).alignment = currency_alignment
        ws1.cell(row=row, column=6, value=item.total_tax_included_price).alignment = currency_alignment
        ws1.cell(row=row, column=7, value=item.total_service_amount).alignment = currency_alignment
        ws1.cell(row=row, column=8, value=item.total_tax).alignment = currency_alignment
        ws1.cell(row=row, column=9, value=item.total_balance).alignment = currency_alignment

    # 写入总计行
    total_row = len(stats_data.groups) + 2
    ws1.cell(row=total_row, column=1, value=stats_data.group_total.manager_name).font = Font(bold=True)
    for col in range(2, 10):
        cell = ws1.cell(row=total_row, column=col)
        cell.font = Font(bold=True)
        cell.alignment = currency_alignment

    ws1.cell(row=total_row, column=2, value=stats_data.group_total.user_count)
    ws1.cell(row=total_row, column=3, value=stats_data.group_total.order_count)
    ws1.cell(row=total_row, column=4, value=stats_data.group_total.product_count)
    ws1.cell(row=total_row, column=5, value=stats_data.group_total.total_internal_price)
    ws1.cell(row=total_row, column=6, value=stats_data.group_total.total_tax_included_price)
    ws1.cell(row=total_row, column=7, value=stats_data.group_total.total_service_amount)
    ws1.cell(row=total_row, column=8, value=stats_data.group_total.total_tax)
    ws1.cell(row=total_row, column=9, value=stats_data.group_total.total_balance)

    # 调整列宽
    ws1.column_dimensions['A'].width = 20
    ws1.column_dimensions['B'].width = 10
    ws1.column_dimensions['C'].width = 12
    ws1.column_dimensions['D'].width = 12
    ws1.column_dimensions['E'].width = 18
    ws1.column_dimensions['F'].width = 18
    ws1.column_dimensions['G'].width = 15
    ws1.column_dimensions['H'].width = 12
    ws1.column_dimensions['I'].width = 15

    # 第二个Sheet：用户详情
    ws2 = wb.create_sheet("用户详情")

    # 写入用户详情表头
    headers2 = ["用户", "订单总数", "商品总数", "订单总团购价格", "订单总含税价格", "总服务价格", "总税额", "总结余"]
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # 按课题组分组
    grouped_users = {}
    for user in stats_data.details:
        group_id = user.group_id
        if group_id not in grouped_users:
            grouped_users[group_id] = {
                "group_name": user.group_name,
                "users": []
            }
        grouped_users[group_id]["users"].append(user)

    # 写入用户详情数据
    current_row = 2
    for group_id in sorted(grouped_users.keys()):
        group = grouped_users[group_id]

        # 写入课题组小计标题
        group_title_cell = ws2.cell(row=current_row, column=1, value=f"{group['group_name']}")
        group_title_cell.fill = group_header_fill
        group_title_cell.font = Font(bold=True)
        ws2.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=8)
        current_row += 1

        # 计算课题组小计
        group_order_count = 0
        group_product_count = 0
        group_internal_price = 0.0
        group_tax_included_price = 0.0
        group_service_amount = 0.0
        group_tax = 0.0
        group_balance = 0.0

        # 写入该课题组的用户
        for user in group["users"]:
            ws2.cell(row=current_row, column=1, value=user.username).alignment = cell_alignment
            ws2.cell(row=current_row, column=2, value=user.order_count).alignment = cell_alignment
            ws2.cell(row=current_row, column=3, value=user.product_count).alignment = cell_alignment
            ws2.cell(row=current_row, column=4, value=user.total_internal_price).alignment = currency_alignment
            ws2.cell(row=current_row, column=5, value=user.total_tax_included_price).alignment = currency_alignment
            ws2.cell(row=current_row, column=6, value=user.total_service_amount).alignment = currency_alignment
            ws2.cell(row=current_row, column=7, value=user.total_tax).alignment = currency_alignment
            ws2.cell(row=current_row, column=8, value=user.total_balance).alignment = currency_alignment

            group_order_count += user.order_count
            group_product_count += user.product_count
            group_internal_price += user.total_internal_price
            group_tax_included_price += user.total_tax_included_price
            group_service_amount += user.total_service_amount
            group_tax += user.total_tax
            group_balance += user.total_balance

            current_row += 1

        # 写入课题组小计
        subtotal_cell = ws2.cell(row=current_row, column=1, value=f"{group['group_name']} 小计")
        subtotal_cell.font = Font(bold=True)
        for col in range(2, 9):
            cell = ws2.cell(row=current_row, column=col)
            cell.font = Font(bold=True)
            cell.alignment = currency_alignment

        ws2.cell(row=current_row, column=2, value=group_order_count)
        ws2.cell(row=current_row, column=3, value=group_product_count)
        ws2.cell(row=current_row, column=4, value=group_internal_price)
        ws2.cell(row=current_row, column=5, value=group_tax_included_price)
        ws2.cell(row=current_row, column=6, value=group_service_amount)
        ws2.cell(row=current_row, column=7, value=group_tax)
        ws2.cell(row=current_row, column=8, value=group_balance)
        current_row += 2  # 空一行

    # 写入总计行
    ws2.cell(row=current_row, column=1, value=stats_data.detail_total.username).font = Font(bold=True)
    for col in range(2, 9):
        cell = ws2.cell(row=current_row, column=col)
        cell.font = Font(bold=True)
        cell.alignment = currency_alignment

    ws2.cell(row=current_row, column=2, value=stats_data.detail_total.order_count)
    ws2.cell(row=current_row, column=3, value=stats_data.detail_total.product_count)
    ws2.cell(row=current_row, column=4, value=stats_data.detail_total.total_internal_price)
    ws2.cell(row=current_row, column=5, value=stats_data.detail_total.total_tax_included_price)
    ws2.cell(row=current_row, column=6, value=stats_data.detail_total.total_service_amount)
    ws2.cell(row=current_row, column=7, value=stats_data.detail_total.total_tax)
    ws2.cell(row=current_row, column=8, value=stats_data.detail_total.total_balance)

    # 调整列宽
    ws2.column_dimensions['A'].width = 20
    ws2.column_dimensions['B'].width = 12
    ws2.column_dimensions['C'].width = 12
    ws2.column_dimensions['D'].width = 18
    ws2.column_dimensions['E'].width = 18
    ws2.column_dimensions['F'].width = 15
    ws2.column_dimensions['G'].width = 12
    ws2.column_dimensions['H'].width = 15

    # 保存到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    logger.info(
        f"按用户统计导出成功: 用户={current_user.username}",
        extra={
            "user_id": current_user.id,
            "user_type": current_user.user_type,
            "start_date": start_date,
            "end_date": end_date
        }
    )

    from urllib.parse import quote
    filename = f"按用户统计_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    encoded_filename = quote(filename, safe='')

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"
        }
    )

